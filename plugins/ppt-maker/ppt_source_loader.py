"""Source loading and lightweight parsing for ppt-maker."""

from __future__ import annotations

import csv
import html
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from xml.etree import ElementTree


class SourceParseError(RuntimeError):
    """Raised when a source cannot be parsed."""


class MissingDependencyError(SourceParseError):
    """Raised when parsing requires an optional dependency group."""

    def __init__(self, dependency_group: str, message: str) -> None:
        super().__init__(message)
        self.dependency_group = dependency_group


@dataclass(slots=True)
class ParsedSource:
    kind: str
    title: str
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)


class SourceLoader:
    """Parse user sources into compact text context."""

    TEXT_EXTS = {".txt", ".md", ".markdown", ".rst", ".log"}
    CSV_EXTS = {".csv", ".tsv"}
    DOC_EXTS = {".pdf", ".docx", ".pptx", ".xlsx"}

    def detect_kind(self, source: str | Path) -> str:
        text = str(source).strip()
        if re.match(r"^https?://", text, flags=re.I):
            return "url"
        suffix = Path(text).suffix.lower()
        if suffix in self.TEXT_EXTS:
            return "markdown" if suffix in {".md", ".markdown"} else "text"
        if suffix in self.CSV_EXTS:
            return "csv"
        if suffix == ".pdf":
            return "pdf"
        if suffix == ".docx":
            return "docx"
        if suffix == ".pptx":
            return "pptx"
        if suffix == ".xlsx":
            return "xlsx"
        return "unknown"

    async def parse(self, source: str | Path, *, kind: str | None = None) -> ParsedSource:
        detected = kind or self.detect_kind(source)
        if detected == "url":
            return await self.parse_url(str(source))
        path = Path(source)
        if not path.exists() or not path.is_file():
            raise SourceParseError(f"Source file not found: {path}")
        if detected in {"text", "markdown"}:
            return self.parse_text(path, kind=detected)
        if detected == "csv":
            return self.parse_csv(path)
        if detected == "pdf":
            return self.parse_pdf(path)
        if detected == "docx":
            return self.parse_docx(path)
        if detected == "pptx":
            return self.parse_pptx(path)
        if detected == "xlsx":
            return self.parse_xlsx(path)
        raise SourceParseError(f"Unsupported source type: {path.suffix or detected}")

    def parse_text(self, path: Path, *, kind: str = "text") -> ParsedSource:
        text = self._read_text(path)
        return ParsedSource(
            kind=kind,
            title=path.stem,
            text=text,
            metadata={"chars": len(text), "path": str(path)},
        )

    def parse_csv(self, path: Path) -> ParsedSource:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        rows: list[list[str]] = []
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            for index, row in enumerate(reader):
                rows.append(row)
                if index >= 50:
                    break
        if not rows:
            raise SourceParseError("CSV file is empty")
        headers = rows[0]
        preview = rows[1:11]
        text_lines = [
            f"Dataset: {path.name}",
            f"Columns ({len(headers)}): {', '.join(headers)}",
            f"Preview rows: {len(preview)}",
        ]
        for row in preview:
            text_lines.append(" | ".join(row))
        return ParsedSource(
            kind="csv",
            title=path.stem,
            text="\n".join(text_lines),
            metadata={
                "columns": headers,
                "preview_row_count": len(preview),
                "path": str(path),
                "delimiter": delimiter,
            },
        )

    def parse_pdf(self, path: Path) -> ParsedSource:
        try:
            from pypdf import PdfReader  # type: ignore[import-not-found]
        except Exception as exc:  # noqa: BLE001
            raise MissingDependencyError(
                "doc_parsing",
                "PDF parsing requires the optional doc_parsing dependency group.",
            ) from exc
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages[:30]]
        text = "\n\n".join(page.strip() for page in pages if page.strip())
        return ParsedSource(
            kind="pdf", title=path.stem, text=text, metadata={"pages": len(reader.pages)}
        )

    def parse_docx(self, path: Path) -> ParsedSource:
        try:
            import docx  # type: ignore[import-not-found]
        except Exception as exc:  # noqa: BLE001
            raise MissingDependencyError(
                "doc_parsing",
                "DOCX parsing requires the optional doc_parsing dependency group.",
            ) from exc
        document = docx.Document(str(path))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        return ParsedSource(
            kind="docx",
            title=path.stem,
            text="\n\n".join(paragraphs),
            metadata={"paragraphs": len(paragraphs)},
        )

    def parse_pptx(self, path: Path) -> ParsedSource:
        texts: list[str] = []
        slide_count = 0
        with zipfile.ZipFile(path) as archive:
            slide_names = sorted(
                name
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            )
            slide_count = len(slide_names)
            for name in slide_names:
                texts.extend(self._extract_pptx_text(archive.read(name)))
        text = "\n".join(item for item in texts if item.strip())
        return ParsedSource(
            kind="pptx",
            title=path.stem,
            text=text,
            metadata={"slides": slide_count, "text_runs": len(texts)},
        )

    def parse_xlsx(self, path: Path) -> ParsedSource:
        try:
            import openpyxl  # type: ignore[import-not-found]
        except Exception as exc:  # noqa: BLE001
            raise MissingDependencyError(
                "table_processing",
                "XLSX parsing requires the optional table_processing dependency group.",
            ) from exc
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        sheet = workbook.active
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            rows.append(["" if cell is None else str(cell) for cell in row])
            if index >= 50:
                break
        headers = rows[0] if rows else []
        preview = rows[1:11]
        text = "\n".join([" | ".join(headers), *(" | ".join(row) for row in preview)])
        return ParsedSource(
            kind="xlsx",
            title=path.stem,
            text=text,
            metadata={
                "sheet": sheet.title,
                "columns": headers,
                "preview_row_count": len(preview),
            },
        )

    async def parse_url(self, url: str) -> ParsedSource:
        try:
            import httpx
        except Exception as exc:  # noqa: BLE001
            raise MissingDependencyError("doc_parsing", "URL parsing requires httpx.") from exc
        async with httpx.AsyncClient(timeout=15, follow_redirects=True) as client:
            response = await client.get(url)
            response.raise_for_status()
        text = self._html_to_text(response.text)
        title = self._extract_html_title(response.text) or url
        return ParsedSource(
            kind="url",
            title=title,
            text=text,
            metadata={"url": url, "status_code": response.status_code, "chars": len(text)},
        )

    def _read_text(self, path: Path) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        raise SourceParseError(f"Unable to decode text file: {path.name}")

    def _extract_pptx_text(self, xml_bytes: bytes) -> list[str]:
        namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        root = ElementTree.fromstring(xml_bytes)
        return [node.text or "" for node in root.findall(".//a:t", namespace)]

    def _html_to_text(self, content: str) -> str:
        content = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", content)
        content = re.sub(r"(?s)<[^>]+>", " ", content)
        content = html.unescape(content)
        return re.sub(r"\s+", " ", content).strip()

    def _extract_html_title(self, content: str) -> str | None:
        match = re.search(r"(?is)<title[^>]*>(.*?)</title>", content)
        if not match:
            return None
        return html.unescape(re.sub(r"\s+", " ", match.group(1))).strip()
